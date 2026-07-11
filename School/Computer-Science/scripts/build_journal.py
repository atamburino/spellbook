"""Build a formatted master's journal DOCX from a Markdown draft.

The output intentionally follows the provided sample journal style:
Letter page, 1-inch margins, Times New Roman 12 pt, double-spaced body,
centered opening block, centered bold section headings.
"""

from __future__ import annotations

import argparse
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


BASE_FONT = "Times New Roman"
BASE_SIZE = Pt(12)
DEFAULT_AUTHOR = "Andy Tamburino"


def slugify(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    return value.strip("-") or "journal"


def today_label() -> str:
    return date.today().strftime("%B %d, %Y").replace(" 0", " ")


def parse_front_matter(text: str) -> tuple[dict[str, str], str]:
    if not text.startswith("---"):
        return {}, text.strip()

    parts = text.split("---", 2)
    if len(parts) < 3:
        return {}, text.strip()

    meta: dict[str, str] = {}
    for raw_line in parts[1].splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or ":" not in line:
            continue
        key, value = line.split(":", 1)
        meta[key.strip().lower()] = value.strip().strip('"')

    return meta, parts[2].strip()


def set_run_font(run, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = BASE_FONT
    run.font.size = BASE_SIZE
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def format_paragraph(paragraph, *, alignment=None, bold=False, space_before=0, space_after=0) -> None:
    if alignment is not None:
        paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 2.0
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    for run in paragraph.runs:
        set_run_font(run, bold=bold)


def add_paragraph(doc: Document, text: str, *, alignment=None, bold=False, before=0, after=0):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)
    format_paragraph(paragraph, alignment=alignment, bold=bold, space_before=before, space_after=after)
    return paragraph


def add_blank_lines(doc: Document, count: int) -> None:
    for _ in range(count):
        paragraph = doc.add_paragraph()
        format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BASE_FONT
    normal.font.size = BASE_SIZE
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)


def add_title_block(doc: Document, meta: dict[str, str]) -> None:
    assignment = meta.get("assignment", "Journal")
    journal_date = meta.get("date", today_label())
    title = meta.get("title", assignment)
    subtitle = meta.get("subtitle", "")

    add_paragraph(doc, assignment, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True, before=36, after=6)
    add_paragraph(doc, journal_date, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_blank_lines(doc, int(meta.get("title_gap_lines", "11")))
    add_paragraph(doc, title, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    if subtitle:
        add_paragraph(doc, subtitle, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=6)


def split_markdown_blocks(body: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    current: list[str] = []

    def flush_paragraph() -> None:
        if current:
            blocks.append(("paragraph", " ".join(line.strip() for line in current).strip()))
            current.clear()

    for raw_line in body.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            flush_paragraph()
            continue
        if line.startswith("## "):
            flush_paragraph()
            blocks.append(("heading", line[3:].strip()))
            continue
        if line.startswith("# "):
            flush_paragraph()
            blocks.append(("heading", line[2:].strip()))
            continue
        current.append(line)

    flush_paragraph()
    return blocks


def add_body(doc: Document, body: str) -> None:
    for kind, text in split_markdown_blocks(body):
        if not text:
            continue
        if kind == "heading":
            add_paragraph(doc, text, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        else:
            add_paragraph(doc, text, alignment=WD_ALIGN_PARAGRAPH.LEFT)


def output_path_for(input_path: Path, meta: dict[str, str]) -> Path:
    if "output" in meta:
        output = Path(meta["output"])
        if not output.is_absolute():
            output = input_path.parent / output
        return output

    assignment = meta.get("assignment", input_path.stem)
    author = meta.get("author", DEFAULT_AUTHOR).split()[-1]
    filename = f"{slugify(assignment)}-{slugify(author)}.docx"
    return input_path.parent.parent / "generated" / filename


def build(input_path: Path, output_override: Path | None = None) -> Path:
    text = input_path.read_text(encoding="utf-8-sig")
    meta, body = parse_front_matter(text)

    doc = Document()
    configure_document(doc)
    add_title_block(doc, meta)
    add_body(doc, body)

    output_path = output_override or output_path_for(input_path, meta)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a formatted journal DOCX from Markdown.")
    parser.add_argument("draft", type=Path, help="Path to the Markdown draft.")
    parser.add_argument("--out", type=Path, help="Optional output DOCX path.")
    args = parser.parse_args()

    output = build(args.draft, args.out)
    print(output)


if __name__ == "__main__":
    main()
